"""Reproduce supplied lifecycle verification and approximate item-overlap summaries.
Run: python reproduce_supporting_analyses.py
The workbook formulas are preserved. Initial and verified codes are reconstructed
from the recorded verdict, and explicit final-code overrides are retained.
This reproduces arithmetic, not an independent reassessment of instrument content.
"""
from pathlib import Path
from collections import Counter
from statistics import median
from openpyxl import load_workbook
from router import route, core_set
from standards import STANDARDS, ROUTES
HERE=Path(__file__).resolve().parent
w=load_workbook(HERE/'data/lifecycle_verification.xlsx',data_only=False)
pairs=[];changes=Counter()
for row in list(w['Coding'].values)[4:]:
    name,cat,stage,stage_name,initial,evidence,verdict,reason,verified,final,anchor=row
    assert verdict in ('Agree','Disagree: should be 1','Disagree: should be 0'), (name,stage,verdict)
    v=initial if verdict=='Agree' else int(verdict.endswith('1'))
    f=int(final) if isinstance(final,(int,float)) else v
    # Final formulas in this supplied workbook retain I(row); numeric overrides implement the item check.
    if isinstance(final,str):
        assert final.startswith('=IF(I'), (name,stage,final)
    assert f==int(stage in STANDARDS[name]['stages']), (name,stage,f)
    pairs.append((initial,f));changes['added']+=int(initial==0 and f==1);changes['removed']+=int(initial==1 and f==0)
n=len(pairs);same=sum(a==b for a,b in pairs);pa=same/n
p0=sum(a for a,b in pairs)/n;p1=sum(b for a,b in pairs)/n;pe=p0*p1+(1-p0)*(1-p1)
print(f'Lifecycle cells {n}; agreement {same}/{n} ({pa:.1%}); kappa {(pa-pe)/(1-pe):.4f}; added {changes["added"]}; removed {changes["removed"]}')
items=list(load_workbook(HERE/'data/item_inventory.xlsx',data_only=True).active.values)[1:]
print(f'Inventory {len(items)} rows; {len({r[0] for r in items})} instruments; {len({r[5] for r in items})} requirement codes')
ns=[];us=[];ols=[]
for typ in [*ROUTES,'protocol']:
    core=core_set(route(is_trial_protocol=True) if typ=='protocol' else route(primary_type=typ))
    selected=[row for row in items if row[0] in core]
    ni=len(selected);nu=len({row[5] for row in selected});ol=100*(1-nu/ni)
    missing=sorted(core-{row[0] for row in items})
    print(f'{typ}: core={len(core)}, nominal={ni}, codes={nu}, overlap={ol:.1f}%, not itemised={missing}')
    if typ!='protocol':ns.append(ni);us.append(nu);ols.append(ol)
print(f'Ten-route medians: items={median(ns)}, codes={median(us)}, overlap={median(ols):.1f}%')
print(f'Ranges: items={min(ns)}-{max(ns)}, codes={min(us)}-{max(us)}, overlap={min(ols):.1f}-{max(ols):.1f}%')
cc=[row for row in items if row[0] in ('FUTURE-AI','STANDING Together')]
print(f'Cross-cutting instruments: items={len(cc)}, codes={len({r[5] for r in cc})}')
